"""Agents' knowledge: files become a searchable knowledge graph, and an agent is briefed
from its own files (and only its own) before it speaks."""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import re
import time

import pytest
from fakes import FakeProvider, default_script
from fastapi.testclient import TestClient

from moot.core import ModelRef, Persona, Round, RoundPlan, RoundType, SessionConfig
from moot.knowledge import KnowledgeService, KnowledgeStore, UnreadableFile, search
from moot.knowledge.chunking import chunk_sections
from moot.knowledge.extraction import entity_key
from moot.knowledge.parsers import Section, read_sections
from moot.providers import ProviderConnectionError
from moot.server import SessionManager, create_app
from moot.storage import export_markdown

MODEL = ModelRef(provider="fake", model="fake-model")
EMBEDDER = ModelRef(provider="fake", model="fake-embed")

OFFICES = "\n".join(
    [
        "# Office options",
        "",
        "Lisbon office rent is about 25 euros per square metre a month. "
        "Lisbon is in the same time zone as London.",
        "",
        "Berlin office rent is about 38 euros per square metre a month. "
        "Berlin has the largest pool of senior engineers in the region.",
        "",
        "Remote work saves the office cost entirely, but the team meets less often.",
    ]
).encode()

GRAPH = {
    "entities": [
        {"name": "Lisbon", "type": "city", "description": "A possible office location."},
        {"name": "Berlin", "type": "city", "description": "A possible office location."},
        {"name": "London", "type": "city", "description": ""},
        {"name": "Remote work", "type": "option", "description": "Working without an office."},
    ],
    "facts": [
        {
            "subject": "Lisbon",
            "relation": "rent",
            "object": "25 euros per m2",
            "statement": "Office rent in Lisbon is about 25 euros per square metre a month.",
        },
        {
            "subject": "Lisbon",
            "relation": "same time zone as",
            "object": "London",
            "statement": "Lisbon is in the same time zone as London.",
        },
        {
            "subject": "Berlin",
            "relation": "rent",
            "object": "38 euros per m2",
            "statement": "Office rent in Berlin is about 38 euros per square metre a month.",
        },
        {
            "subject": "Berlin",
            "relation": "has",
            "object": "senior engineers",
            "statement": "Berlin has the largest pool of senior engineers in the region.",
        },
        {
            "subject": "Remote work",
            "relation": "saves",
            "object": "office cost",
            "statement": "Remote work saves the office cost entirely.",
        },
        {
            "subject": "Nobody",
            "relation": "x",
            "object": "y",
            "statement": "Dropped: unknown subject.",
        },
    ],
}


def embed_text(text: str, dims: int = 1024) -> list[float]:
    """A bag of hashed words: texts sharing words get similar vectors."""
    vector = [0.0] * dims
    for word in re.findall(r"\w+", text.lower()):
        vector[int(hashlib.md5(word.encode()).hexdigest(), 16) % dims] += 1.0
    return vector


def is_extraction(request) -> bool:
    return request.messages[0].content.startswith("You build a knowledge graph")


class KnowledgeProvider(FakeProvider):
    """Answers extraction requests with GRAPH and can embed; everything else as usual."""

    def __init__(self, script=None, *, embed_error=None, extract_error=None):
        def answer(request):
            if is_extraction(request):
                return extract_error if extract_error else [json.dumps(GRAPH)]
            return (script or default_script)(request)

        super().__init__(answer)
        self.embed_error = embed_error
        self.embedded: list[str] = []

    async def embed(self, model, texts):
        if self.embed_error:
            raise self.embed_error
        self.embedded += texts
        return [embed_text(t) for t in texts]


def service_with(provider) -> KnowledgeService:
    return KnowledgeService(KnowledgeStore(":memory:"), lambda preset: provider)


def index(
    service: KnowledgeService,
    data: bytes = OFFICES,
    name: str = "offices.md",
    embedding=EMBEDDER,
    model=MODEL,
):
    async def go():
        base = service.create_base("Ines", embedding)
        document = service.add_document(base.id, name, data, model)
        await service.wait()
        return base.id, document.id

    return asyncio.run(go())


def tiny_pdf(pages: list[str]) -> bytes:
    """A minimal PDF with one line of text per page (Helvetica)."""
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        None,
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode()}\nendstream")
        content = len(objects)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content} 0 R "
            "/Resources << /Font << /F1 3 0 R >> >> >>"
        )
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


# ------------------------------------------------------------------ reading and splitting


def test_reading_text_markdown_word_and_pdf():
    assert read_sections("notes.txt", "Café prices\n\n\nare up".encode()) == [
        Section("Café prices\n\nare up")
    ]
    from docx import Document

    document = Document()
    document.add_paragraph("Lisbon has good weather.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Rent"
    table.rows[0].cells[1].text = "25 euros"
    buffer = io.BytesIO()
    document.save(buffer)
    [word] = read_sections("plan.docx", buffer.getvalue())
    assert "Lisbon has good weather." in word.text and "Rent | 25 euros" in word.text

    pages = read_sections("report.pdf", tiny_pdf(["Rent in Lisbon is low", "Berlin has engineers"]))
    assert [(s.page, s.text) for s in pages] == [
        (1, "Rent in Lisbon is low"),
        (2, "Berlin has engineers"),
    ]


def test_files_moot_cant_read_are_refused_clearly():
    with pytest.raises(UnreadableFile, match="Moot reads"):
        read_sections("slides.pptx", b"x")
    with pytest.raises(UnreadableFile, match="no text"):
        read_sections("scan.pdf", tiny_pdf([""]))
    with pytest.raises(UnreadableFile, match="could not be read as a PDF"):
        read_sections("broken.pdf", b"not a pdf")


def test_passages_are_small_overlapping_and_stay_on_their_page():
    text = " ".join(f"Sentence number {i} says something about offices." for i in range(80))
    chunks = chunk_sections(
        [Section(text, page=3), Section("Short page.", page=4)], size=400, overlap=60
    )
    assert len(chunks) > 5
    assert all(len(c.text) <= 460 for c in chunks)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    assert chunks[-1].page == 4 and chunks[-1].text == "Short page."
    assert all(c.page == 3 for c in chunks[:-1])
    first_end = chunks[0].text[-30:].split()[-1]
    assert first_end in chunks[1].text  # the overlap carries over


# ------------------------------------------------------------------ indexing


def test_a_file_becomes_a_knowledge_graph_with_vectors():
    provider = KnowledgeProvider()
    service = service_with(provider)
    base_id, document_id = index(service)
    base = service.store.base(base_id)
    [document] = base.documents
    assert (document.status, document.stage, document.error, document.notes) == (
        "ready",
        "",
        None,
        [],
    )
    assert document.chunks_total == base.passages == 1 and document.chunks_mapped == 1
    assert base.entities == 4 and base.facts == 5  # the fact about an unknown subject was dropped
    assert base.embedding == EMBEDDER
    assert service.store.embedding_of(base_id)[1] == 1024
    entities, facts = service.store.graph(base_id)
    assert entities[0]["name"] in ("Lisbon", "Berlin")  # the most connected first
    assert any(f["object"] == "London" for f in facts)
    assert any(text.startswith("Lisbon: A possible office") for text in provider.embedded)


def test_the_brief_finds_facts_and_passages_by_words_and_meaning():
    service = service_with(KnowledgeProvider())
    base_id, _ = index(service)
    brief = asyncio.run(service.brief(base_id, "How much does an office in Lisbon cost?"))
    statements = [f.statement for f in brief.facts]
    assert statements[0] == "Office rent in Lisbon is about 25 euros per square metre a month."
    assert "Lisbon is in the same time zone as London." in statements
    assert brief.passages and brief.passages[0].source == "offices.md"
    assert brief.sources == ["offices.md"]
    text = brief.text()
    assert text.startswith("Your own knowledge (from your files; only you can see this):")
    assert "[offices.md]" in text and "Never cite a source that isn't listed here" in text
    nothing = asyncio.run(service.brief(base_id, "zebra quantum"))
    assert nothing.empty and nothing.text() == ""


def test_without_an_embedding_model_search_uses_words_and_the_graph():
    service = service_with(KnowledgeProvider())
    base_id, _ = index(service, embedding=None)
    brief = search(service.store, base_id, "Is Berlin good for hiring senior engineers?")
    assert (
        brief.facts[0].statement == "Berlin has the largest pool of senior engineers in the region."
    )


def test_pdf_passages_are_cited_by_page():
    service = service_with(KnowledgeProvider())
    base_id, _ = index(service, tiny_pdf(["Intro page", "Lisbon office rent is low"]), "report.pdf")
    brief = asyncio.run(service.brief(base_id, "Lisbon rent"))
    assert "report.pdf p.2" in brief.sources


def test_failures_keep_the_file_searchable_with_a_note():
    down = ProviderConnectionError("Fake", "can't reach it")
    service = service_with(KnowledgeProvider(embed_error=down, extract_error=down))
    base_id, document_id = index(service)
    document = service.store.document(document_id)
    assert document.status == "ready"
    assert any("Embedding the passages failed" in n for n in document.notes)
    assert document.chunks_failed == 1 and any("couldn't be mapped" in n for n in document.notes)
    assert search(service.store, base_id, "Lisbon rent").passages  # still found by its words

    _, bad = index(service, b"\x00\x01", "scan.pdf")
    failed = service.store.document(bad)
    assert failed.status == "failed" and "could not be read" in failed.error


def test_mapping_stops_after_failures_in_a_row():
    text = "\n\n".join(f"Paragraph {i}. " + "words " * 200 for i in range(8)).encode()
    service = service_with(KnowledgeProvider(extract_error=ProviderConnectionError("Fake", "down")))
    service._concurrency = 1
    _, document_id = index(service, text, "long.txt")
    document = service.store.document(document_id)
    assert document.chunks_total > 5 and document.chunks_failed == 5
    assert any("stopped after 5 failed calls" in n for n in document.notes)


def test_removing_a_file_removes_what_was_learned_from_it():
    service = service_with(KnowledgeProvider())
    base_id, document_id = index(service)
    extra = asyncio.run(
        _add(service, base_id, b"Porto is a cheaper option than Lisbon.", "more.txt")
    )
    asyncio.run(service.delete_document(base_id, document_id))
    base = service.store.base(base_id)
    assert [d.id for d in base.documents] == [extra]
    assert base.passages == 1
    # entities still mentioned by the other file stay; the rest go with their facts
    assert base.entities == 4 and base.facts == 5
    asyncio.run(service.delete_document(base_id, extra))
    base = service.store.base(base_id)
    assert (base.passages, base.entities, base.facts) == (0, 0, 0)
    assert search(service.store, base_id, "Lisbon rent").empty


async def _add(service, base_id, data, name):
    document = service.add_document(base_id, name, data, MODEL)
    await service.wait()
    return document.id


def test_files_cut_off_by_a_restart_are_marked_failed():
    store = KnowledgeStore(":memory:")
    base = store.create_base("x")
    document = store.add_document(base.id, "a.txt", 3)
    KnowledgeService(store)
    assert store.document(document.id).status == "failed"
    assert "Moot stopped" in store.document(document.id).error


def test_entity_names_match_loosely():
    assert entity_key("The Lisbon Office") == entity_key("lisbon office") == "lisbon office"
    assert entity_key("Café") == "cafe"


# ------------------------------------------------------------------ in a discussion


def test_an_agent_is_briefed_from_its_own_files_only():
    provider = KnowledgeProvider()

    async def scenario():
        manager = SessionManager(provider_factory=lambda preset: provider)
        base = manager.knowledge.create_base("Ava", EMBEDDER)
        manager.knowledge.add_document(base.id, "offices.md", OFFICES, MODEL)
        await manager.knowledge.wait()
        config = SessionConfig(
            topic="Should we open an office in Lisbon?",
            agents=[Persona(name="Ava", knowledge=base.id), Persona(name="Ben")],
            round_plan=RoundPlan(rounds=[Round(type=RoundType.OPENING)]),
            default_model=MODEL,
        )
        record = manager.create(config)
        manager.start(record.id)
        await record.task
        await manager.shutdown()
        return record.state

    state = asyncio.run(scenario())
    turns = {
        r.messages[0].content.split(",")[0]: r.messages[-1].content
        for r in provider.requests
        if r.messages[0].content.startswith("You are ") and "Instruction:" in r.messages[-1].content
    }
    ava = next(v for k, v in turns.items() if k.startswith("You are Ava"))
    ben = next(v for k, v in turns.items() if k.startswith("You are Ben"))
    assert "Your own knowledge" in ava and "Office rent in Lisbon" in ava
    assert "Your own knowledge" not in ben
    ava_message = next(m for m in state.messages if m.speaker_id == state.config.agents[0].id)
    assert ava_message.sources == ["offices.md"]
    assert "*Briefed from its files: offices.md*" in export_markdown(state)


def test_a_session_with_missing_knowledge_is_refused():
    manager = SessionManager(provider_factory=lambda preset: FakeProvider())
    config = SessionConfig(
        topic="x",
        agents=[Persona(name="Ava", knowledge="kb_gone"), Persona(name="Ben")],
        default_model=MODEL,
    )
    with pytest.raises(ValueError, match="knowledge files no longer exist"):
        manager.create(config)


# ------------------------------------------------------------------ the API


def test_the_api_adds_indexes_searches_and_removes_files():
    provider = KnowledgeProvider()
    app = create_app(SessionManager(provider_factory=lambda preset: provider))
    with TestClient(app, base_url="http://127.0.0.1") as client:
        base = client.post(
            "/api/knowledge", json={"name": "Ines", "embedding": EMBEDDER.model_dump()}
        ).json()
        added = client.put(
            f"/api/knowledge/{base['id']}/files?name=offices.md&provider=fake&model=fake-model",
            content=OFFICES,
            headers={"Content-Type": "application/octet-stream"},
        )
        assert added.status_code == 202 and added.json()["status"] in (
            "queued",
            "indexing",
            "ready",
        )
        deadline = time.monotonic() + 10
        while (
            client.get(f"/api/knowledge/{base['id']}").json()["documents"][0]["status"] != "ready"
        ):
            assert time.monotonic() < deadline
            time.sleep(0.05)
        found = client.get(
            f"/api/knowledge/{base['id']}/search", params={"q": "Lisbon rent"}
        ).json()
        assert found["facts"][0]["source"] == "offices.md"
        assert found["briefing"].startswith("Your own knowledge")
        graph = client.get(f"/api/knowledge/{base['id']}/graph").json()
        assert len(graph["entities"]) == 4
        assert [b["id"] for b in client.get("/api/knowledge").json()] == [base["id"]]

        refused = client.put(f"/api/knowledge/{base['id']}/files?name=deck.pptx", content=b"x")
        assert refused.status_code == 415 and "Moot reads" in refused.json()["detail"]
        assert (
            client.put("/api/knowledge/kb_none/files?name=a.txt", content=b"x").status_code == 404
        )

        document_id = added.json()["id"]
        assert client.delete(f"/api/knowledge/{base['id']}/files/{document_id}").status_code == 204
        assert client.get(f"/api/knowledge/{base['id']}").json()["documents"] == []
        assert client.delete(f"/api/knowledge/{base['id']}").status_code == 204
        assert client.get(f"/api/knowledge/{base['id']}").status_code == 404
