"""Reading an agent's files into plain text: PDF (pypdf), Word (python-docx), text and Markdown.

Each file becomes sections; a PDF keeps its page numbers, so passages can be cited by page.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import PurePath

SUPPORTED = {
    ".pdf": "PDF",
    ".docx": "Word",
    ".txt": "text",
    ".md": "Markdown",
    ".markdown": "Markdown",
}
MAX_FILE_BYTES = 25 * 1024 * 1024


class UnreadableFile(ValueError):
    """The file can't be read: wrong type, damaged, protected, or without any text."""


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None = None  # 1-based, for PDFs


def kind_of(filename: str) -> str:
    suffix = PurePath(filename).suffix.lower()
    if suffix not in SUPPORTED:
        allowed = ", ".join(sorted({s.lstrip(".").upper() for s in SUPPORTED} - {"MARKDOWN"}))
        raise UnreadableFile(f"{filename}: Moot reads {allowed} files.")
    return suffix


def read_sections(filename: str, data: bytes) -> list[Section]:
    """The file's text, in order. Raises UnreadableFile."""
    suffix = kind_of(filename)
    if len(data) > MAX_FILE_BYTES:
        raise UnreadableFile(f"{filename} is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    if suffix == ".pdf":
        sections = _pdf(filename, data)
    elif suffix == ".docx":
        sections = _docx(filename, data)
    else:
        sections = [Section(_decode(data))]
    sections = [Section(_tidy(s.text), s.page) for s in sections]
    sections = [s for s in sections if s.text]
    if not sections:
        raise UnreadableFile(
            f"{filename} has no text Moot can read (a scanned PDF needs text recognition first)."
        )
    return sections


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(encoding)
            if encoding == "utf-16" and data[:2] not in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return text
        except UnicodeDecodeError:
            continue
    return data.decode("cp1252", errors="replace")


def _pdf(filename: str, data: bytes) -> list[Section]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001 - any failure means it stays locked
                raise UnreadableFile(f"{filename} is password-protected.") from exc
        return [
            Section(page.extract_text() or "", page=number)
            for number, page in enumerate(reader.pages, start=1)
        ]
    except UnreadableFile:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise UnreadableFile(f"{filename} could not be read as a PDF ({exc}).") from exc


def _docx(filename: str, data: bytes) -> list[Section]:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises several kinds
        raise UnreadableFile(f"{filename} could not be read as a Word file.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(c for c in cells if c))
    return [Section("\n".join(parts))]


def _tidy(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\r", "\n").split("\n")]
    out: list[str] = []
    for line in lines:
        if line or (out and out[-1]):
            out.append(line)
    return "\n".join(out).strip()
